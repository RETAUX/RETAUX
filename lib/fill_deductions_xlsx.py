"""
Rellena el template Tax Deductions Individual con datos del cliente.
Uso: python3 fill_deductions_xlsx.py <data.json> <output.xlsx>
"""
import sys, json, datetime, warnings, traceback
warnings.filterwarnings('ignore')
try:
    from openpyxl import load_workbook
except ImportError as e:
    print(f"ERROR: {e}", file=sys.stderr)
    sys.exit(1)

DATA_JSON = sys.argv[1]
OUT_XLSX  = sys.argv[2]
TEMPLATE  = sys.argv[3] if len(sys.argv) > 3 else None

with open(DATA_JSON) as f:
    d = json.load(f)

wb = load_workbook(TEMPLATE or OUT_XLSX, keep_vba=False)

def write(ws, row, col, val):
    cell = ws.cell(row, col)
    if hasattr(cell, 'value'):
        cell.value = val

# ── Client Details ───────────────────────────────────────────────────────────
cd = wb['Client Details']
name = d.get('fullName', '')
write(cd, 3, 2, name.upper() if name else '')

# Year ending: use FY end date
fy_year = d.get('taxYear', '2025-26')
fy_end_year = int(fy_year.split('-')[1]) if '-' in fy_year else 2026
write(cd, 3, 4, datetime.datetime(fy_end_year, 6, 30))

write(cd, 8, 2, d.get('address', ''))
write(cd, 12, 3, d.get('occupation', ''))

dob_str = d.get('dateOfBirth', '') or d.get('birthDate', '')
if dob_str:
    try:
        write(cd, 13, 3, datetime.datetime.fromisoformat(dob_str[:10]))
    except Exception:
        write(cd, 13, 3, dob_str)

# Spouse
has_spouse = bool(d.get('spouseName') or d.get('hasSpouse'))
write(cd, 14, 3, 'Yes' if has_spouse else 'No')

# Medicare / PHI
write(cd, 15, 3, 'No')  # Medicare exemption — default No unless specified
write(cd, 16, 3, 'Yes' if d.get('hasPhi') else 'No')
write(cd, 17, 3, 'Not Provided')

# Deduction category checkboxes (rows 21-29)
ded = d.get('ded', {})
has_d1 = (ded.get('D1', 0) or 0) > 0
has_d2 = (ded.get('D2', 0) or 0) > 0
has_d3 = (ded.get('D3', 0) or 0) > 0
has_d4 = (ded.get('D4', 0) or 0) > 0
has_d5 = (ded.get('D5', 0) or 0) > 0
has_d9 = (ded.get('D9', 0) or 0) > 0
has_d10 = (ded.get('D10', 0) or 0) > 0

write(cd, 21, 3, 'Yes' if has_d1 else 'No')   # Motor Vehicle
write(cd, 22, 3, 'Yes' if has_d2 else 'No')   # Work-related travel
write(cd, 23, 3, 'No')                          # Home Office (not tracked yet)
write(cd, 24, 3, 'Yes' if has_d3 else 'No')   # Clothing
write(cd, 25, 3, 'Yes' if has_d4 else 'No')   # Self-education
write(cd, 26, 3, 'No')                          # Assets/depreciation
write(cd, 27, 3, 'Yes' if has_d10 else 'No')  # Tax agent fees
write(cd, 28, 3, 'Yes' if has_d9 else 'No')   # Donations
write(cd, 29, 3, 'No')                          # Super contributions

# ── D1. Work-related Car Expenses (cents per km) ────────────────────────────
if has_d1 and d.get('d1_km'):
    d1 = wb['D1. Work-related Car Expenses']
    write(d1, 18, 1, d['d1_km'])
    write(d1, 18, 2, d.get('d1_rate', 0.88))
    write(d1, 18, 4, d.get('d1_ref', 'Work-related motor vehicle travel'))

# ── D2. Work-related travel ──────────────────────────────────────────────────
if has_d2:
    d2 = wb['D2. Work-related travel expense']
    write(d2, 4, 1, 'Work-related travel expenses')
    write(d2, 4, 3, ded['D2'])
    write(d2, 4, 4, 'Calculated from uploaded receipts')

# ── D3. Work-related clothing ────────────────────────────────────────────────
if has_d3:
    d3 = wb['D3. Work-related clothing']
    write(d3, 4, 1, 'Work-related clothing/uniform')
    write(d3, 4, 3, ded['D3'])
    write(d3, 4, 4, 'Calculated from uploaded receipts')

# ── D4. Self-education ───────────────────────────────────────────────────────
if has_d4:
    d4 = wb['D4. Self-education']
    write(d4, 4, 1, 'Self-education / Training')
    write(d4, 4, 3, ded['D4'])
    write(d4, 4, 4, 'Calculated from uploaded receipts')

# ── D5. Other work-related expenses ─────────────────────────────────────────
if has_d5:
    d5 = wb['D5. Other work-related expenses']
    write(d5, 4, 1, 'Other work-related expenses (phone/tech/other)')
    write(d5, 4, 3, ded['D5'])
    write(d5, 4, 4, 'Calculated from uploaded receipts')

# ── D9. Gifts or donations ───────────────────────────────────────────────────
if has_d9:
    d9 = wb['D9. Gifts or donations']
    write(d9, 4, 1, 'Donations')
    write(d9, 4, 3, ded['D9'])
    write(d9, 4, 4, 'Calculated from uploaded receipts')

# ── D10. Cost of managing tax affairs (RETAUX fee) ──────────────────────────
if has_d10:
    d10 = wb['D10. Cost of managing tax affai']
    write(d10, 4, 1, 'Tax agent fees (RETAUX)')
    write(d10, 4, 3, ded['D10'])
    write(d10, 4, 4, 'Tax agent fee – RETAUX (deductible in the following FY)')

wb.save(OUT_XLSX)
print('OK')
