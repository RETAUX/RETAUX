"""
Actualiza los Travel Log de FY 2025-26 para que el total de km (B14)
coincida exactamente con lo que el cliente declaró en RETAUX (vehicles[].km).

Lógica de distribución:
  - base_oneway = total_km // 112   (2 legs × 56 días)
  - Las primeras N días usan (base_oneway + 1) ida/vuelta
  - El resto usa base_oneway ida/vuelta
  - Esto garantiza: sum(E+H) == total_km exacto
  - Si no hay km en RETAUX: archivo queda con km=0 en todas las filas

Uso: python3 scripts/update-travel-km.py
"""

import openpyxl
from pathlib import Path

OUT_DIR = Path(__file__).parent.parent / 'exports' / 'TravelLogs'

# ── Km declarados en RETAUX (FY 2025-26, vehicles[].km sumado) ───────────────
# Nombre exacto del cliente (debe coincidir con el campo 'name' en fill-travel-log.py)
RETAUX_KM = {
    'Mark Wust':                       5000,
    'Luciano Guillermo Barbieri':      4980,
    'Teo David Meglio':                4990,
    'Maria Agustina González Roulet':  3500,
    'Melany Borda':                    4990,
    'Marina Pulichino':                4990,
    'Hugo Toca Sainz-Rozas':           5000,
    'Javier Alberto Pedrotti':         4780,
    'Felipe Giraudo':                  4000,
    'Joaquin vandersluis':             4200,
    'Florencia Cristina Rey':          4900,
    'Agustina Saenz Batista':          4890,
    'Giuliano Oberrti':                4990,
    'Daniel Romero Garcia':            4890,
    'Ramiro Fera':                     4990,
}

# Clientes en el batch sin km en RETAUX → archivo con km=0
# (los que no están en RETAUX_KM)

WORKDAYS = 56  # Apr 13 – Jun 30 2026, excluyendo Queen's Birthday VIC (Jun 8)


def km_distribution(total_km: int, workdays: int = WORKDAYS):
    """
    Devuelve lista de (outbound_km, return_km) por día (len == workdays)
    tal que sum(out + ret) == total_km exactamente.

    total_km siempre es par (RETAUX acepta enteros y los valores son pares),
    por lo que cada día tiene outbound == return.
    """
    legs = workdays * 2          # 112 legs total
    base = total_km // legs      # km por leg para la mayoría
    extra = total_km % legs      # legs adicionales que llevan (base + 1)
    # extra siempre par porque total_km y legs son ambos pares
    extra_days = extra // 2      # días con (base+1) ida Y (base+1) vuelta

    result = []
    for i in range(workdays):
        km = base + 1 if i < extra_days else base
        result.append((km, km))
    return result


def update_file(path: Path, total_km: int):
    wb = openpyxl.load_workbook(path)

    # ── Summary: actualizar B9 al km-oneway típico ──────────────────────────
    ws = wb['Summary']
    if total_km > 0:
        typical_oneway = round(total_km / (WORKDAYS * 2))
        ws['B9'] = typical_oneway
    else:
        ws['B9'] = 0

    # ── Travel Log: actualizar E y H de cada fila con fecha ─────────────────
    tl = wb['Travel Log']
    dist = km_distribution(total_km) if total_km > 0 else [(0, 0)] * WORKDAYS

    day_idx = 0
    for row_idx in range(6, 62):
        if tl.cell(row_idx, 1).value is None:
            continue
        if day_idx >= len(dist):
            break
        out_km, ret_km = dist[day_idx]
        tl.cell(row_idx, 5).value = out_km   # E: outbound km
        tl.cell(row_idx, 8).value = ret_km   # H: return km
        day_idx += 1

    wb.save(path)
    return typical_oneway if total_km > 0 else 0


# ── Procesar todos los archivos del batch ────────────────────────────────────
files = sorted(f for f in OUT_DIR.glob('*.xlsx') if not f.name.startswith('~$'))

ok = 0
blanked = 0
errors = []

for f in files:
    # Extraer nombre del cliente del nombre del archivo
    client_name = f.name.replace(' FY2025-26_Travel Log.xlsx', '')

    total_km = RETAUX_KM.get(client_name, 0)

    try:
        oneway = update_file(f, total_km)
        if total_km > 0:
            print(f'  ✓  {client_name:45s} → {total_km} km  (B9={oneway})')
            ok += 1
        else:
            print(f'  –  {client_name:45s} → sin datos (km=0)')
            blanked += 1
    except Exception as e:
        errors.append((client_name, str(e)))
        print(f'  ✗  {client_name}: {e}')

print(f'\n=== RESUMEN ===')
print(f'Actualizados con km de RETAUX: {ok}')
print(f'Sin datos (km=0):              {blanked}')
print(f'Errores:                       {len(errors)}')
if errors:
    for name, err in errors:
        print(f'  ! {name}: {err}')

# ── Verificar que los totales sean correctos ─────────────────────────────────
print('\n── Verificación de totales ──')
for name, expected_km in RETAUX_KM.items():
    safe = name.replace('/', '_').replace(':', '_')
    path = OUT_DIR / f'{safe} FY2025-26_Travel Log.xlsx'
    if not path.exists():
        print(f'  ? {name}: archivo no encontrado')
        continue
    wb = openpyxl.load_workbook(path, data_only=True)
    tl = wb['Travel Log']
    total = 0
    for row_idx in range(6, 62):
        e = tl.cell(row_idx, 5).value or 0
        h = tl.cell(row_idx, 8).value or 0
        total += e + h
    match = '✓' if total == expected_km else f'✗ got {total}'
    print(f'  {match}  {name}: {expected_km} km')
